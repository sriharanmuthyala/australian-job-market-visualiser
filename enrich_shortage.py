#!/usr/bin/env python3
"""
Merge the Jobs and Skills Australia 2025 Unit Group Shortage List into data.json.

Usage:
    pip install openpyxl
    python enrich_shortage.py data.json "2025 Unit Group Shortage List - 4 digit ANZSCO.xlsx"

The script matches 4-digit ANZSCO unit-group codes and writes:
- shortage: true/false
- shortage_rating: the original national rating, when available
- shortage_driver: the original shortage-driver field, when available
"""

import argparse
import json
import re
from pathlib import Path

from openpyxl import load_workbook


def norm(value):
    return re.sub(r"\s+", " ", str(value or "").strip()).lower()


def code4(value):
    m = re.search(r"\b(\d{4})\b", str(value or ""))
    return m.group(1) if m else None


def find_header(ws):
    for r in range(1, min(ws.max_row, 25) + 1):
        vals = [norm(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
        joined = " | ".join(vals)
        if ("anzsco" in joined or "unit group" in joined or "occupation code" in joined) and "shortage" in joined:
            return r, vals
    raise RuntimeError("Could not identify the header row in the shortage workbook.")


def choose_col(headers, include, exclude=()):
    for i, h in enumerate(headers, start=1):
        if all(x in h for x in include) and not any(x in h for x in exclude):
            return i
    return None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("data", nargs="?", default="data.json")
    ap.add_argument("workbook")
    args = ap.parse_args()

    wb = load_workbook(args.workbook, data_only=True)
    ws = wb[wb.sheetnames[0]]
    header_row, headers = find_header(ws)

    code_col = (
        choose_col(headers, ("anzsco",), ("6 digit",))
        or choose_col(headers, ("unit group", "code"))
        or choose_col(headers, ("occupation", "code"))
    )
    rating_col = (
        choose_col(headers, ("national", "shortage"))
        or choose_col(headers, ("australia", "shortage"))
        or choose_col(headers, ("shortage", "rating"), ("state",))
    )
    driver_col = choose_col(headers, ("shortage", "driver"))

    if not code_col or not rating_col:
        raise RuntimeError(
            f"Could not find code/rating columns. Headers were: {headers}"
        )

    ratings = {}
    for r in range(header_row + 1, ws.max_row + 1):
        code = code4(ws.cell(r, code_col).value)
        if not code:
            continue
        rating_raw = ws.cell(r, rating_col).value
        rating = str(rating_raw).strip() if rating_raw is not None else ""
        driver = str(ws.cell(r, driver_col).value).strip() if driver_col and ws.cell(r, driver_col).value is not None else None
        if rating:
            ratings[code] = {"rating": rating, "driver": driver}

    path = Path(args.data)
    doc = json.loads(path.read_text(encoding="utf-8"))
    matched = 0
    in_shortage = 0

    for d in doc["occupations"]:
        rec = ratings.get(str(d.get("code", "")))
        if not rec:
            continue
        rating = rec["rating"]
        low = rating.lower()
        is_shortage = "shortage" in low and "no shortage" not in low
        d["shortage"] = is_shortage
        d["shortage_rating"] = rating
        if rec["driver"]:
            d["shortage_driver"] = rec["driver"]
        matched += 1
        in_shortage += int(is_shortage)

    meta = doc.setdefault("meta", {})
    meta["shortage_source"] = "Jobs and Skills Australia, 2025 Occupation Shortage List"
    meta["shortage_level"] = "4-digit ANZSCO unit groups"
    meta["shortage_status"] = "official"

    path.write_text(json.dumps(doc, indent=1, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"Matched {matched} occupations; {in_shortage} rated as shortage in some form.")


if __name__ == "__main__":
    main()
