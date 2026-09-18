#!/usr/bin/env python3
"""
Populate pay_weekly in data.json from the official Jobs and Skills Australia
ANZSCO occupation profiles workbook.

The profile workbook reports median weekly earnings for ANZSCO 4-digit
occupations. JSA defines this as median pay for full-time, non-managerial
employees paid at adult rates, sourced from the ABS Survey of Employee
Earnings and Hours.

Usage:
    python enrich_pay.py data.json
"""

import argparse
import html
import json
import re
import tempfile
import urllib.parse
import urllib.request
from pathlib import Path

LANDING = "https://www.jobsandskills.gov.au/data/occupation-and-industry-profiles"
FALLBACK_XLSX = (
    "https://www.jobsandskills.gov.au/sites/default/files/2026-07/"
    "ANZSCO%20Occupation%20data%20-%20February%202026.xlsx"
)
UA = "Mozilla/5.0 (compatible; ausjobs-treemap/1.0; +https://github.com/sriharanmuthyala/australian-job-market-visualiser)"


def request(url):
    return urllib.request.Request(
        url,
        headers={
            "User-Agent": UA,
            "Referer": LANDING,
            "Accept": "*/*",
        },
    )


def discover_workbook_url():
    """Find the current ANZSCO occupation workbook link; fall back if markup changes."""
    try:
        with urllib.request.urlopen(request(LANDING), timeout=30) as response:
            page = response.read().decode("utf-8", "replace")
        links = re.findall(r'href=["\']([^"\']+\.xlsx(?:\?[^"\']*)?)', page, flags=re.I)
        links = [html.unescape(x) for x in links]
        for href in links:
            decoded = urllib.parse.unquote(href).lower()
            if "anzsco" in decoded and "occupation" in decoded:
                return urllib.parse.urljoin(LANDING, href)
    except Exception as exc:
        print(f"Could not discover workbook link ({exc}); using known JSA URL.")
    return FALLBACK_XLSX


def download(url):
    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    tmp.close()
    print(f"Downloading JSA occupation profiles:\n  {url}")
    with urllib.request.urlopen(request(url), timeout=90) as response, open(tmp.name, "wb") as out:
        out.write(response.read())
    return tmp.name


def norm(value):
    return re.sub(r"\s+", " ", str(value if value is not None else "")).strip().lower()


def pay_number(value):
    if value is None:
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return int(round(value))
    text = str(value).strip()
    if not text or text.upper() in {"N/A", "NA", "-", "—"}:
        return None
    digits = re.sub(r"[^0-9.]", "", text)
    if not digits:
        return None
    return int(round(float(digits)))


def find_overview_sheet(workbook):
    preferred = workbook["Table_1"] if "Table_1" in workbook.sheetnames else None
    candidates = [preferred] if preferred else []
    candidates += [ws for ws in workbook.worksheets if ws is not preferred]

    for ws in candidates:
        for row_num, row in enumerate(ws.iter_rows(min_row=1, max_row=20, values_only=True), 1):
            headers = [norm(v) for v in row]
            code_idx = next((i for i, h in enumerate(headers) if "anzsco" in h and "code" in h), None)
            pay_idx = next((i for i, h in enumerate(headers) if "median" in h and "weekly" in h and "earn" in h), None)
            if code_idx is not None and pay_idx is not None:
                return ws, row_num, code_idx, pay_idx
    raise RuntimeError("Could not find ANZSCO Code / Median weekly earnings columns in JSA workbook.")


def load_pay(path):
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    ws, header_row, code_idx, pay_idx = find_overview_sheet(wb)
    values = {}

    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if code_idx >= len(row):
            continue
        raw_code = row[code_idx]
        if raw_code is None:
            continue
        if isinstance(raw_code, (int, float)) and not isinstance(raw_code, bool):
            code = str(int(raw_code))
        else:
            code = re.sub(r"\.0$", "", str(raw_code).strip())
        if not (code.isdigit() and len(code) == 4):
            continue
        pay = pay_number(row[pay_idx] if pay_idx < len(row) else None)
        values[code] = pay

    wb.close()
    return values, ws.title


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("data", nargs="?", default="data.json")
    args = parser.parse_args()

    data_path = Path(args.data)
    doc = json.loads(data_path.read_text(encoding="utf-8"))

    workbook_url = discover_workbook_url()
    workbook_path = download(workbook_url)
    pay, sheet_name = load_pay(workbook_path)

    matched = 0
    with_pay = 0
    for occupation in doc.get("occupations", []):
        code = str(occupation.get("code", ""))
        if code in pay:
            matched += 1
            occupation["pay_weekly"] = pay[code]
            if pay[code] is not None:
                with_pay += 1

    meta = doc.setdefault("meta", {})
    meta["pay_source"] = "Jobs and Skills Australia occupation profiles; ABS Survey of Employee Earnings and Hours"
    meta["pay_reference_period"] = "May 2025"
    meta["pay_profile_snapshot"] = "February 2026"
    meta["pay_definition"] = "Median weekly earnings for full-time non-managerial employees paid at adult rates."
    meta["pay_source_url"] = LANDING
    meta["pay_workbook"] = workbook_url
    meta["pay_sheet"] = sheet_name

    data_path.write_text(json.dumps(doc, indent=1, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"Matched {matched} four-digit occupations; populated pay for {with_pay}.")


if __name__ == "__main__":
    main()
